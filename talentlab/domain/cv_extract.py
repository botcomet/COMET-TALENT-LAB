"""Extraction et structuration d'un CV (§21).

Principes :
- un échec d'extraction est une *erreur explicite* : jamais de score inventé (test 14) ;
- les dates ne sont jamais inventées : une date absente reste inconnue ;
- le texte conservé est le texte source — toute preuve est une tranche exacte de ce texte.
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date
from typing import Any

from .text import fold, normalize_text

MAX_BYTES_DEFAULT = 10 * 1024 * 1024
MAX_PAGES_DEFAULT = 40
MIN_CHARS_DEFAULT = 250
MAX_CHARS_DEFAULT = 150_000          # un CV de 40 pages ≈ 120 000 caractères ; au-delà, ce n'est pas un CV (et le temps d'analyse doit rester borné)


class ExtractionError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass
class Extraction:
    text: str
    pages: int
    quality: str                     # "ok" | "partial"
    warnings: list[str] = field(default_factory=list)
    page_offsets: list[tuple[int, int]] = field(default_factory=list)
    mime: str = "text/plain"

    def page_of(self, pos: int) -> int | None:
        for i, (a, b) in enumerate(self.page_offsets, start=1):
            if a <= pos < b:
                return i
        return None


# ----------------------------------------------------------------- détection / extraction
def sniff(data: bytes) -> str:
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            return "corrupt"
        return "unsupported"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "doc"
    head = data[:4000]
    if b"\x00" not in head:
        try:
            head.decode("utf-8")
            return "txt"
        except UnicodeDecodeError:
            try:
                head.decode("latin-1")
                printable = sum(1 for b in head if 32 <= b < 127 or b in (9, 10, 13) or b >= 160)
                if printable / max(1, len(head)) > 0.95:
                    return "txt"
            except Exception:  # pragma: no cover
                pass
    return "unsupported"


def _alpha_ratio(text: str) -> float:
    non_ws = [c for c in text if not c.isspace()]
    if not non_ws:
        return 0.0
    return sum(1 for c in non_ws if c.isalnum()) / len(non_ws)


def _garbled(text: str) -> bool:
    if not text:
        return True
    if text.count("�") / max(1, len(text)) > 0.02:
        return True
    toks = text.split()
    if len(toks) > 40 and sum(1 for t in toks if len(t) == 1) / len(toks) > 0.45:
        return True            # texte espacé lettre à lettre / polices mal encodées
    return _alpha_ratio(text) < 0.55


def _pdf(data: bytes, max_pages: int) -> tuple[list[str], list[str]]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
    warnings: list[str] = []
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
    except (PdfReadError, ValueError, KeyError, OSError, AssertionError) as e:
        raise ExtractionError("corrupt", "Le fichier PDF est corrompu ou mal formé et ne peut pas être lu.") from e
    except Exception as e:  # noqa: BLE001 — pypdf peut lever des exceptions variées sur des fichiers hostiles
        raise ExtractionError("corrupt", "Le fichier PDF est illisible.") from e
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")
        except Exception:  # noqa: BLE001
            ok = 0
        if not ok:
            raise ExtractionError("encrypted", "Le PDF est protégé par un mot de passe : le fournir ou importer une version déprotégée.")
    n = len(reader.pages)
    if n > max_pages:
        raise ExtractionError("too_many_pages", f"Le document compte {n} pages (maximum {max_pages}) : probablement pas un CV.")
    pages: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001
            pages.append("")
            warnings.append(f"Page {i} illisible : son contenu est ignoré.")
    if sum(len(p.strip()) for p in pages) < 40:     # essai de repli avec un autre extracteur
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                alt = [(pg.extract_text() or "") for pg in pdf.pages[:max_pages]]
            if sum(len(p.strip()) for p in alt) > sum(len(p.strip()) for p in pages):
                pages = alt
        except Exception:  # noqa: BLE001
            pass
    return pages, warnings


def _docx(data: bytes) -> list[str]:
    try:
        import docx
        d = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise ExtractionError("corrupt", "Le fichier Word est corrompu et ne peut pas être lu.") from e
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append("  ".join(c.text.strip() for c in row.cells if c.text.strip()))
    return ["\n".join(parts)]


def extract_text(data: bytes, filename: str = "", *, max_bytes: int = MAX_BYTES_DEFAULT, max_pages: int = MAX_PAGES_DEFAULT,
                 min_chars: int = MIN_CHARS_DEFAULT, max_chars: int = MAX_CHARS_DEFAULT) -> Extraction:
    if not data:
        raise ExtractionError("empty", "Le fichier est vide.")
    if len(data) > max_bytes:
        raise ExtractionError("too_large", f"Le fichier dépasse {max_bytes // (1024 * 1024)} Mo.")
    kind = sniff(data)
    if kind in ("unsupported", "doc"):
        raise ExtractionError("unsupported_type", "Format non pris en charge (PDF, DOCX ou texte uniquement)." if kind == "unsupported"
                              else "Le format Word 97-2003 (.doc) n'est pas pris en charge : l'enregistrer en .docx ou PDF.")
    if kind == "corrupt":
        raise ExtractionError("corrupt", "Archive Word corrompue.")
    warnings: list[str] = []
    mime = {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "txt": "text/plain"}[kind]
    if kind == "pdf":
        pages, warnings = _pdf(data, max_pages)
    elif kind == "docx":
        pages = _docx(data)
    else:
        try:
            pages = [data.decode("utf-8")]
        except UnicodeDecodeError:
            pages = [data.decode("latin-1")]
    cleaned = [normalize_text(p) for p in pages]
    total = sum(len(p) for p in cleaned)
    if total == 0 or (kind == "pdf" and total < 40):
        raise ExtractionError("no_text_layer", "Aucun texte exploitable : PDF probablement scanné (image). L'OCR n'est pas disponible en V1 : importer une version textuelle.")
    if _garbled("\n".join(cleaned)):
        raise ExtractionError("garbled", "Le texte extrait est illisible (polices mal encodées ou caractères corrompus) : aucun score ne peut être calculé de façon fiable.")
    if total > max_chars:
        raise ExtractionError("too_long", f"Le texte extrait compte {total} caractères (maximum {max_chars}) : probablement pas un CV.")
    if total < min_chars:
        raise ExtractionError("insufficient_content", f"Contenu insuffisant ({total} caractères) pour une analyse fiable d'un CV.")
    quality = "ok"
    avg = total / max(1, len(cleaned))
    if kind == "pdf" and len(cleaned) > 1 and avg < 400:
        quality = "partial"
        warnings.append("Peu de texte par page : une partie du document est peut-être une image non lue.")
    if any("illisible" in w for w in warnings):
        quality = "partial"
    text = ""
    offsets: list[tuple[int, int]] = []
    for p in cleaned:
        start = len(text)
        text += p
        offsets.append((start, len(text)))
        text += "\n\n"
    return Extraction(text=text.strip("\n"), pages=len(cleaned), quality=quality, warnings=warnings, page_offsets=offsets, mime=mime)


# ----------------------------------------------------------------- dates
_MONTHS = {
    "jan": 1, "janv": 1, "janvier": 1, "january": 1, "feb": 2, "fev": 2, "fevr": 2, "fevrier": 2, "february": 2,
    "mar": 3, "mars": 3, "march": 3, "avr": 4, "apr": 4, "avril": 4, "april": 4, "mai": 5, "may": 5,
    "juin": 6, "jun": 6, "june": 6, "juil": 7, "jul": 7, "juillet": 7, "july": 7, "aout": 8, "aug": 8, "august": 8,
    "sep": 9, "sept": 9, "septembre": 9, "september": 9, "oct": 10, "octobre": 10, "october": 10,
    "nov": 11, "novembre": 11, "november": 11, "dec": 12, "decembre": 12, "december": 12,
}
_MON = "|".join(sorted(_MONTHS, key=len, reverse=True))
_DATE = rf"(?:\d{{1,2}}\s*[/.]\s*\d{{4}}|\d{{4}}\s*/\s*\d{{1,2}}(?!\d)|(?:(?:{_MON})\.?\s+)?\d{{4}})"
_END = rf"(?:{_DATE}|present|aujourd'?hui|actuel(?:lement)?|ce jour|en cours|current|now|date|maintenant|today|ongoing)"
_RANGE = re.compile(rf"(?P<s>{_DATE})\s*(?:-|–|—|a|au|to|jusqu'?a|->|→)\s*(?P<e>{_END})", re.I)
_RANGE_SINCE = re.compile(rf"(?:depuis|since|from)\s+(?P<s>{_DATE})", re.I)


@dataclass(frozen=True)
class YM:
    year: int
    month: int | None             # None = mois inconnu

    def ordinal(self, default_month: int) -> int:
        return self.year * 12 + (self.month or default_month) - 1

    def iso(self) -> str:
        return f"{self.year:04d}-{self.month:02d}" if self.month else f"{self.year:04d}"


def _parse_ym(tok: str) -> YM | None:
    t = fold(tok).strip()
    if m := re.fullmatch(r"(\d{1,2})\s*[/.]\s*(\d{4})", t):
        mo = int(m.group(1))
        return YM(int(m.group(2)), mo if 1 <= mo <= 12 else None)
    if m := re.fullmatch(r"(\d{4})\s*/\s*(\d{1,2})", t):
        mo = int(m.group(2))
        return YM(int(m.group(1)), mo if 1 <= mo <= 12 else None)
    if m := re.fullmatch(rf"(?:({_MON})\.?\s+)?(\d{{4}})", t):
        return YM(int(m.group(2)), _MONTHS.get(m.group(1)) if m.group(1) else None)
    return None


def _is_open(tok: str) -> bool:
    return bool(re.fullmatch(r"present|aujourd'?hui|actuel(?:lement)?|ce jour|en cours|current|now|date|maintenant|today|ongoing", fold(tok).strip()))


_QUANTITY_AFTER = re.compile(r"\s*(?:k\b|m\b|%|€|\$|euros?|eur\b|messages?|msgs?|transactions?|requetes?|requests?|utilisateurs?|users?|jours?|mois|ans?\b|tps|ko\b|mo\b|go\b|to\b|lignes?|brokers?|topics?|serveurs?)", re.I)
MIN_PLAUSIBLE_YEAR = 1960


def plausible_range(start_tok: str, end_tok: str, line_after: str, today: date) -> bool:
    """Une plage de dates d'expérience : années dans [1960 ; année courante + 1], jamais une quantité (« 5000 - 8000 messages », « 1500 - 2000 euros »)."""
    sy, ey = _parse_ym(start_tok), (None if _is_open(end_tok) else _parse_ym(end_tok))
    if sy is None or not (MIN_PLAUSIBLE_YEAR <= sy.year <= today.year + 1):
        return False
    if ey is not None and not (MIN_PLAUSIBLE_YEAR <= ey.year <= today.year + 1):
        return False
    return not _QUANTITY_AFTER.match(line_after)


# ----------------------------------------------------------------- structure
_HEADINGS = {
    "skills": r"competences?(?: techniques| cles| principales)?|skills|technologies|expertises?|connaissances techniques|stack(?: technique)?|outils|savoir[- ]faire|environnements? techniques? global|domaines? de competences?",
    "experience": r"experiences?(?: professionnelles?)?|parcours(?: professionnel)?|work experience|professional experience|historique professionnel|missions? realisees|projets? realises",
    "education": r"formations?|education|diplomes?|etudes|cursus",
    "languages": r"langues?|languages?",
    "certifications": r"certifications?|habilitations?|certificats?",
    "summary": r"profil|resume|summary|a propos|objectif|presentation",
    "interests": r"centres? d'interets?|loisirs|hobbies|interests|divers",
}
_HEAD_RX = {k: re.compile(rf"^(?:{v})$") for k, v in _HEADINGS.items()}
_ENV_LINE = re.compile(r"^\s*(?:environnements?(?:\s+techniques?)?|stack(?:\s+technique)?|technologies|outils|tech|environment)\s*[:：]", re.I)
_NON_PRO = re.compile(r"familial|personnel|a titre perso|side project|projet perso|association|benevole|projet etudiant|projet scolaire|travaux pratiques|projet de fin d'etudes|hobby|loisir")


def _heading_of(line: str) -> str | None:
    f = fold(line).strip(" :：-–—#*_\t.")
    if not f or len(f) > 60:
        return None
    for name, rx in _HEAD_RX.items():
        if rx.match(f):
            return name
    return None


@dataclass
class Experience:
    idx: int
    header: str
    company: str
    title: str
    start: str | None
    end: str | None
    is_current: bool
    months: int | None
    precision: str                      # "month" | "year" | "unknown"
    span: tuple[int, int]               # étendue dans le texte
    body_span: tuple[int, int]
    env_spans: list[tuple[int, int]] = field(default_factory=list)
    dates_unknown: bool = False

    def to_dict(self) -> dict[str, Any]:
        return {k: (list(v) if isinstance(v, tuple) else v) for k, v in self.__dict__.items()}


@dataclass
class ParsedCV:
    text: str
    sections: dict[str, list[tuple[int, int]]]
    experiences: list[Experience]
    languages: dict[str, str]
    declared_years: int | None
    computed_years: float | None
    flags: list[str]
    extraction_quality: str = "ok"

    def section_of(self, pos: int) -> str | None:
        for name, spans in self.sections.items():
            for a, b in spans:
                if a <= pos < b:
                    return name
        return None

    def experience_at(self, pos: int) -> Experience | None:
        for e in self.experiences:
            if e.span[0] <= pos < e.span[1]:
                return e
        return None


_LEVELS = {"courant": "courant", "fluent": "courant", "bilingue": "bilingue", "bilingual": "bilingue", "natif": "bilingue", "native": "bilingue",
           "professionnel": "professionnel", "professional": "professionnel", "operationnel": "professionnel", "intermediaire": "intermédiaire",
           "intermediate": "intermédiaire", "scolaire": "scolaire", "notions": "notions", "basic": "notions", "debutant": "notions",
           "b1": "intermédiaire", "b2": "professionnel", "c1": "courant", "c2": "bilingue", "a2": "notions", "a1": "notions"}


def _languages(text: str) -> dict[str, str]:
    f = fold(text)
    out: dict[str, str] = {}
    for lang, rx in (("anglais", r"(?:anglais|english)"), ("allemand", r"(?:allemand|german|deutsch)"), ("espagnol", r"(?:espagnol|spanish)"),
                     ("francais", r"(?:francais|french)")):
        m = re.search(rf"{rx}\s*[:\-–(]*\s*(?:niveau\s*)?({'|'.join(_LEVELS)})\b", f) or re.search(rf"\b({'|'.join(_LEVELS)})\b\s*(?:en\s+)?{rx}", f)
        if m:
            out[lang] = _LEVELS[m.group(1)]
        elif re.search(rf"\b{rx}\b", f) and lang != "francais":
            out.setdefault(lang, "niveau non précisé")
    return out


def _split_header(header: str) -> tuple[str, str]:
    h = header
    for m in sorted(_RANGE.finditer(fold(header)), key=lambda m: -m.start()):     # texte replié : mêmes offsets que l'original
        h = h[:m.start()] + " " + h[m.end():]
    h = re.sub(r"\(\s*\)", " ", h)
    h = re.sub(r"[()\[\]]", " ", h)
    parts = [p.strip(" ,;:-–—|·") for p in re.split(r"\s+[-–—|·@]\s+|\s+chez\s+|\s*\|\s*", h) if p.strip(" ,;:-–—|·")]
    if len(parts) >= 2:
        return parts[0], parts[1]
    return (parts[0] if parts else ""), ""


def parse_cv(text: str, *, today: date | None = None, extraction_quality: str = "ok") -> ParsedCV:
    today = today or date.today()
    lines: list[tuple[int, int, str]] = []
    pos = 0
    for raw in text.split("\n"):
        lines.append((pos, pos + len(raw), raw))
        pos += len(raw) + 1

    # --- sections
    marks: list[tuple[int, str]] = []
    for a, _b, raw in lines:
        h = _heading_of(raw)
        if h:
            marks.append((a, h))
    sections: dict[str, list[tuple[int, int]]] = {}
    for i, (a, name) in enumerate(marks):
        b = marks[i + 1][0] if i + 1 < len(marks) else len(text)
        sections.setdefault(name, []).append((a, b))
    head_end = marks[0][0] if marks else 0
    if head_end > 0:
        sections.setdefault("summary", []).insert(0, (0, head_end))

    # --- expériences : chaque plage de dates dans la zone d'expérience ouvre un bloc
    zones = sections.get("experience") or [(0, len(text))]
    skip = [s for n in ("skills", "education", "languages", "certifications", "interests") for s in sections.get(n, [])]
    cand: list[tuple[int, int, str, str, str | None]] = []   # (line_start, line_end, raw, s_tok, e_tok)
    flags: list[str] = []
    for a, b, raw in lines:
        if not any(z0 <= a < z1 for z0, z1 in zones):
            continue
        if any(s0 <= a < s1 for s0, s1 in skip):
            continue
        f = fold(raw)
        m = next((mm for mm in _RANGE.finditer(f) if plausible_range(mm.group("s"), mm.group("e"), f[mm.end():mm.end() + 24], today)), None)
        if m is None and _RANGE.search(f) and len(raw) <= 160:
            first = _RANGE.search(f)
            if (re.search(r"\b(?:19|20)\d{2}\b", first.group(0)) and not _QUANTITY_AFTER.match(f[first.end():first.end() + 24])):
                flags.append(f"Plage de dates invraisemblable ignorée : « {raw.strip()[:70]} » (années hors de {MIN_PLAUSIBLE_YEAR}–{today.year + 1}).")
        if m and len(raw) <= 160:
            cand.append((a, b, raw, m.group("s"), m.group("e")))
        elif (ms := _RANGE_SINCE.search(f)) and len(raw) <= 160 and plausible_range(ms.group("s"), "present", "", today):
            cand.append((a, b, raw, ms.group("s"), "present"))
    if cand:
        kept: list[tuple[int, int, str, str, str | None]] = []
        last: tuple[int, int] | None = None
        for c in cand:
            sy_, ey_ = _parse_ym(c[3]), (None if _is_open(c[4] or "") else _parse_ym(c[4] or ""))
            rng = (sy_.ordinal(1), (today.year * 12 + today.month - 1) if _is_open(c[4] or "") else (ey_.ordinal(12) if ey_ else -1)) if sy_ else None
            if (last and rng and re.match(r"^\s*[-•*·▪►→]", c[2]) and rng[0] >= last[0] and 0 <= rng[1] <= last[1]):
                continue                          # « - Migration (2019 - 2020) » dans une expérience 2018-2022 : période de projet
            kept.append(c)
            if rng and rng[1] >= 0 and not re.match(r"^\s*[-•*·▪►→]", c[2]):
                last = rng
        cand = kept
    exps: list[Experience] = []
    idx_of = {a: i for i, (a, _b, _r) in enumerate(lines)}
    for k, (a, b, raw, s_tok, e_tok) in enumerate(cand):
        i = idx_of[a]
        h0 = i
        for back in (1, 2):                      # lignes d'en-tête précédentes (société / intitulé)
            j = i - back
            if j < 0:
                break
            la, lb, lraw = lines[j]
            f = fold(lraw)
            if (lraw.strip() and len(lraw) <= 100 and not lraw.strip().endswith((".", ";")) and not re.match(r"^\s*[-•*·▪►→]", lraw)
                    and not _RANGE.search(f) and not _heading_of(lraw) and not _ENV_LINE.match(lraw)):
                h0 = j
            else:
                break
        start_char = lines[h0][0]
        end_char = cand[k + 1][0] if k + 1 < len(cand) else min([z1 for z0, z1 in zones if z0 <= a < z1] + [len(text)])
        if k + 1 < len(cand):         # remonte avant les en-têtes de l'expérience suivante
            ni = idx_of[cand[k + 1][0]]
            nh = ni
            for back in (1, 2):
                j = ni - back
                if j <= i:
                    break
                la, lb, lraw = lines[j]
                if (lraw.strip() and len(lraw) <= 100 and not lraw.strip().endswith((".", ";")) and not re.match(r"^\s*[-•*·▪►→]", lraw)
                        and not _RANGE.search(fold(lraw)) and not _heading_of(lraw) and not _ENV_LINE.match(lraw)):
                    nh = j
                else:
                    break
            end_char = lines[nh][0]
        sy, ey = _parse_ym(s_tok), (None if _is_open(e_tok or "") else _parse_ym(e_tok or ""))
        is_cur = _is_open(e_tok or "")
        months, prec = None, "unknown"
        if sy and (ey or is_cur):
            e_ord = (today.year * 12 + today.month - 1) if is_cur else ey.ordinal(7)   # type: ignore[union-attr]
            s_ord = sy.ordinal(7)
            if e_ord < s_ord:
                flags.append(f"Dates incohérentes (fin avant début) dans « {lines[i][2].strip()[:60]} » : durée non calculée.")
            else:
                months = max(1, e_ord - s_ord + (1 if (sy.month and (ey is None or ey.month)) else 0))
                prec = "month" if (sy.month and (is_cur or (ey and ey.month))) else "year"
        header = "\n".join(lines[x][2].strip() for x in range(h0, i + 1) if lines[x][2].strip())
        company, title = _split_header(header.replace("\n", " — "))
        env = [(la, lb) for la, lb, lraw in lines if start_char <= la < end_char and _ENV_LINE.match(lraw)]
        exps.append(Experience(
            idx=len(exps), header=header, company=company, title=title,
            start=sy.iso() if sy else None, end=(None if is_cur else (ey.iso() if ey else None)), is_current=is_cur,
            months=months, precision=prec, span=(start_char, end_char), body_span=(b, end_char), env_spans=env,
            dates_unknown=months is None,
        ))
    # texte d'expérience hors de toute période datée : une expérience SANS DATES (jamais ignorée ni datée par supposition)
    covered = sorted((e.span for e in exps))
    for z0, z1 in zones:
        cur = z0
        for a, b in [c for c in covered if c[1] > z0 and c[0] < z1] + [(z1, z1)]:
            if a - cur > 120 and text[cur:a].strip():
                chunk = text[cur:a]
                first = next((ln.strip() for ln in chunk.split("\n") if ln.strip() and not _heading_of(ln)), "(expérience non datée)")
                if len(re.sub(r"\s+", "", chunk)) > 100 and not _heading_of(first):
                    company, title = _split_header(first)
                    exps.append(Experience(idx=-1, header=first, company=company, title=title, start=None, end=None, is_current=False, months=None,
                                           precision="unknown", span=(cur, a), body_span=(cur, a),
                                           env_spans=[(la, lb) for la, lb, lraw in lines if cur <= la < a and _ENV_LINE.match(lraw)], dates_unknown=True))
            cur = max(cur, b)
    exps.sort(key=lambda e: e.span[0])
    for i, e in enumerate(exps):
        e.idx = i
    if not exps and zones:
        z0, z1 = zones[0]
        exps.append(Experience(idx=0, header="(expériences non datées)", company="", title="", start=None, end=None, is_current=False,
                               months=None, precision="unknown", span=(z0, z1), body_span=(z0, z1), dates_unknown=True))
        flags.append("Aucune expérience datée n'a pu être identifiée : durées et récence non évaluables (aucune date n'est inventée).")

    # --- cohérence des dates (§15.7)
    intervals = sorted(((_ord(e.start, 7), _end_ord(e, today)) for e in exps if e.months is not None and e.start), key=lambda t: t[0])
    merged: list[list[int]] = []
    for s, e in intervals:
        if merged and s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    computed = round(sum(e - s for s, e in merged) / 12, 1) if merged else None
    declared = None
    if m := re.search(r"(\d{1,2})\s*(?:\+\s*)?(?:ans|annees|years)\s*d['’ ]?\s*(?:experience|exp\b)", fold(text[:2500])):
        declared = int(m.group(1))
    if declared and computed is not None and abs(declared - computed) > 2:
        flags.append(f"Écart de dates : le CV annonce {declared} ans d'expérience, les périodes datées en totalisent {computed} ans. Privilégier les dates précises vérifiées.")
    undated = [e for e in exps if e.dates_unknown and e.header != "(expériences non datées)"]
    if undated:
        flags.append(f"{len(undated)} expérience(s) sans dates exploitables : durée et récence inconnues.")
    return ParsedCV(text=text, sections=sections, experiences=exps, languages=_languages(text), declared_years=declared,
                    computed_years=computed, flags=flags, extraction_quality=extraction_quality)


def _ord(iso: str | None, default_month: int) -> int:
    y, _, m = (iso or "0000").partition("-")
    return int(y) * 12 + (int(m) if m else default_month) - 1


def _end_ord(e: Experience, today: date) -> int:
    if e.is_current:
        return today.year * 12 + today.month - 1
    return _ord(e.end, 7)
